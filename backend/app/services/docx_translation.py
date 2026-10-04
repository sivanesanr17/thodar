from __future__ import annotations

import re
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from docx.document import Document as DocumentObject
    from docx.text.paragraph import Paragraph

from app.services.translation_service import TranslationProvider, TranslationUnavailable

URL_PATTERN = re.compile(r"(?:https?://|www\.)[^\s<>]+", re.IGNORECASE)
NUMBER_ONLY_PATTERN = re.compile(r"\s*(?:\d+|[ivxlcdm]+)\s*", re.IGNORECASE)
CODE_LIKE_PATTERN = re.compile(
    r"^(?:```|#include\b|(?:def|class|import|from|const|let|var|function)\s|<[/!?\w]|\{[\s\S]*\}|\[[\s\S]*\])"
)


def _table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested_table in cell.tables:
                yield from _table_paragraphs(nested_table)


def iter_document_paragraphs(document: DocumentObject):
    """Walk body, table, header and footer paragraphs once, including nested tables."""
    candidates = list(document.paragraphs)
    for table in document.tables:
        candidates.extend(_table_paragraphs(table))

    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            candidates.extend(part.paragraphs)
            for table in part.tables:
                candidates.extend(_table_paragraphs(table))

    seen = set()
    for paragraph in candidates:
        identity = id(paragraph._p)
        if identity not in seen:
            seen.add(identity)
            yield paragraph


def _can_translate(paragraph: Paragraph, text: str) -> bool:
    if not text.strip() or NUMBER_ONLY_PATTERN.fullmatch(text):
        return False
    if URL_PATTERN.fullmatch(text.strip()):
        return False
    if paragraph._p.xpath(".//w:instrText"):
        return False
    if CODE_LIKE_PATTERN.search(text.strip()):
        return False
    return True


def _protect_urls(text: str) -> tuple[str, dict[str, str]]:
    replacements: dict[str, str] = {}

    def replace(match: re.Match[str]) -> str:
        token = f"THODARURLTOKEN{len(replacements)}END"
        replacements[token] = match.group(0)
        return token

    return URL_PATTERN.sub(replace, text), replacements


def append_translated_text(paragraph: Paragraph, text: str) -> None:
    """Append translated text as Unicode without imposing a legacy font encoding."""
    paragraph.add_run(text)


def _replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    text_runs = [run for run in paragraph.runs if run.text]
    if not text_runs:
        run = paragraph.add_run()
        run.text = text
        return

    if len(text_runs) == 1:
        text_runs[0].text = text
        return

    weights = [len(run.text) for run in text_runs]
    total_weight = sum(weights)
    boundaries = [0]
    cumulative_weight = 0

    for weight in weights[:-1]:
        cumulative_weight += weight
        target = round(len(text) * cumulative_weight / total_weight)
        possible_boundaries = [
            index + 1
            for index, character in enumerate(text)
            if character.isspace() and boundaries[-1] < index + 1 < len(text)
        ]
        if possible_boundaries:
            target = min(possible_boundaries, key=lambda boundary: abs(boundary - target))
        target = max(boundaries[-1], min(target, len(text)))
        boundaries.append(target)

    boundaries.append(len(text))
    for run, start, end in zip(text_runs, boundaries, boundaries[1:]):
        run.text = text[start:end]


async def translate_docx(document: DocumentObject, provider: TranslationProvider,
                         source_language: str, target_language: str) -> int:
    """Translate text blocks once per distinct value, preserving document paragraphs."""
    paragraphs_to_translate = []
    for paragraph in iter_document_paragraphs(document):
        source_text = paragraph.text
        if _can_translate(paragraph, source_text):
            paragraphs_to_translate.append((paragraph, source_text))

    unique_texts = list(dict.fromkeys(text for _, text in paragraphs_to_translate))
    prepared = {text: _protect_urls(text) for text in unique_texts}
    cache: dict[str, str] = {}
    translate_many = getattr(provider, "translate_many", None)

    if callable(translate_many):
        batch_size = 8
        for start in range(0, len(unique_texts), batch_size):
            batch = unique_texts[start : start + batch_size]
            protected_texts = [prepared[text][0] for text in batch]
            results = await translate_many(protected_texts, source_language, target_language)
            if len(results) != len(batch):
                raise TranslationUnavailable("The translation service returned an incomplete batch.")
            for original, translated in zip(batch, results):
                for token, url in prepared[original][1].items():
                    translated = translated.replace(token, url)
                cache[original] = translated
    else:
        for original in unique_texts:
            protected_text, urls = prepared[original]
            translated = await provider.translate(protected_text, source_language, target_language)
            for token, url in urls.items():
                translated = translated.replace(token, url)
            cache[original] = translated

    for paragraph, source_text in paragraphs_to_translate:
        _replace_paragraph_text(paragraph, cache[source_text])

    return len(paragraphs_to_translate)
