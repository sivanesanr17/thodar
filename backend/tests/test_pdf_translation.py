import asyncio
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch


try:
    import pymupdf as fitz
    from fastapi.testclient import TestClient
    from app.services.pdf_translation import translate_pdf_to_file
    from app.services.pdf_layout import extract_pdf_layout
    from app.main import app
except ImportError:
    fitz = None
    TestClient = None
    translate_pdf_to_file = None
    extract_pdf_layout = None
    app = None


class RecordingProvider:
    def __init__(self):
        self.calls = []

    async def translate(self, text, source_language, target_language):
        self.calls.append((text, source_language, target_language))
        return f"[{target_language}] {text}"


class FixedTranslationProvider(RecordingProvider):
    async def translate(self, text, source_language, target_language):
        self.calls.append((text, source_language, target_language))
        return "Translated paragraph."


class TamilTranslationProvider(RecordingProvider):
    async def translate(self, text, source_language, target_language):
        self.calls.append((text, source_language, target_language))
        return "தமிழ்நாடு அரசு"


@unittest.skipIf(fitz is None, "PyMuPDF is required for PDF translation tests")
class PdfTranslationTests(unittest.TestCase):
    def test_tamil_paragraphs_use_bamini_and_keep_line_breaks(self):
        from app.models.document import DocumentElement, DocumentPage, StructuredDocument
        from app.services.pdf_docx_renderer import render_translated_docx

        layout = StructuredDocument(source_format="pdf", pages=[DocumentPage(
            page_number=1, width=420, height=620, elements=[DocumentElement(
                type="paragraph", original_text="Tamil text", translated_text="அரசு\nதமிழ்",
                bbox=(20, 30, 180, 80), page=1,
            )],
        )])
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "translated.docx"
            render_translated_docx(layout, output_path, "ta")
            from docx import Document
            output = Document(output_path)
        paragraph = output.paragraphs[0]
        self.assertIn("\n", paragraph.text)
        self.assertIn("<w:br", paragraph._p.xml)
        self.assertTrue(any(run.font.name == "Bamini" for run in paragraph.runs))

    def test_two_column_reading_order_is_column_first(self):
        from app.models.document import DocumentElement, DocumentPage
        from app.services.pdf_docx_renderer import _ordered_page_elements

        elements = [
            DocumentElement(type="paragraph", original_text=text, bbox=bbox, page=1)
            for text, bbox in [
                ("left top", (40, 80, 250, 110)),
                ("right top", (350, 85, 560, 115)),
                ("left bottom", (40, 150, 250, 180)),
                ("right bottom", (350, 155, 560, 185)),
            ]
        ]
        page = DocumentPage(page_number=1, width=600, height=800, elements=elements)

        self.assertEqual(
            [element.original_text for element in _ordered_page_elements(page)],
            ["left top", "left bottom", "right top", "right bottom"],
        )

    def test_translates_repeated_text_into_flowing_docx_paragraphs(self):
        document = fitz.open()
        first = document.new_page(width=420, height=620)
        first.insert_text((50, 70), "A short paragraph for translation.")
        second = document.new_page(width=595, height=842)
        second.insert_text((55, 80), "A short paragraph for translation.")
        provider = FixedTranslationProvider()

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "translated.docx"
            translated_count = asyncio.run(
                translate_pdf_to_file(document, provider, "ta", "en", output_path)
            )
            from docx import Document
            output = Document(output_path)
            self.assertEqual(translated_count, 2)
            self.assertEqual(len(provider.calls), 1)
            self.assertEqual(len(output.sections), 1)
            self.assertEqual([p.text for p in output.paragraphs], [
                "Translated paragraph.", "Translated paragraph."
            ])
            self.assertEqual(output.paragraphs[1].paragraph_format.space_before.pt, 12)
            self.assertFalse(output.paragraphs[0]._p.xpath(".//w:drawing"))
            self.assertFalse(output.sections[0].header._element.xpath(".//w:pict"))
        document.close()

    def test_rejects_image_only_pdf(self):
        document = fitz.open()
        document.new_page()
        provider = RecordingProvider()
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "translated.docx"
            with self.assertRaisesRegex(ValueError, "no_extractable_text"):
                asyncio.run(translate_pdf_to_file(document, provider, "ta", "en", output_path))
        self.assertEqual(provider.calls, [])
        document.close()

    def test_extracts_pdf_text_geometry_and_font_metadata(self):
        document = fitz.open()
        page = document.new_page(width=420, height=620)
        page.insert_text((50, 70), "Positioned heading", fontsize=20, fontname="hebo")

        structured = extract_pdf_layout(document)

        self.assertEqual(structured.source_format, "pdf")
        self.assertEqual((structured.pages[0].width, structured.pages[0].height), (420, 620))
        element = structured.pages[0].elements[0]
        self.assertEqual(element.original_text, "Positioned heading")
        self.assertEqual(element.type, "heading")
        self.assertEqual(element.page, 1)
        self.assertGreater(element.bbox[0], 0)
        self.assertEqual(element.font_size, 20)
        self.assertTrue(element.font_name)
        document.close()

    def test_pdf_endpoint_returns_docx_with_bamini_tamil_text(self):
        source = fitz.open()
        source.new_page(width=420, height=620).insert_text(
            (50, 70), "A short paragraph for translation."
        )
        pdf_bytes = source.tobytes()
        source.close()

        with patch("app.routes.translations.create_translation_provider", return_value=TamilTranslationProvider()):
            response = TestClient(app).post(
                "/translate/pdf",
                data={"source_language": "en", "target_language": "ta"},
                files={"file": ("source.pdf", pdf_bytes, "application/pdf")},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("wordprocessingml.document", response.headers["content-type"])
        self.assertIn("source_translated.docx", response.headers["content-disposition"])
        from docx import Document
        translated = Document(io.BytesIO(response.content))
        self.assertIn("jkpo;ehL muR", translated.paragraphs[0].text)
        self.assertTrue(any(run.font.name == "Bamini" for run in translated.paragraphs[0].runs))
        self.assertFalse(translated.paragraphs[0]._p.xpath(".//w:drawing"))


if __name__ == "__main__":
    unittest.main()
