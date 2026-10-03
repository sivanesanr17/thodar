import unittest

from docx import Document

from app.services.bamini import split_bamini_segments, unicode_to_bamini
from app.services.docx_translation import _replace_paragraph_text


class BaminiConversionTests(unittest.TestCase):
    def test_encodes_tamil_phrase(self):
        self.assertEqual(unicode_to_bamini("தமிழ்நாடு அரசு"), "jkpo;ehL muR")

    def test_keeps_english_separate_from_bamini_font_run(self):
        document = Document()
        paragraph = document.add_paragraph()
        source_run = paragraph.add_run("source")
        source_run.bold = True

        _replace_paragraph_text(paragraph, "தமிழ்நாடு Government", "ta")

        self.assertEqual(paragraph.text, "jkpo;ehL Government")
        self.assertEqual(paragraph.runs[0].font.name, "Bamini")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertEqual(paragraph.runs[1].text, " Government")
        self.assertTrue(paragraph.runs[1].bold)

    def test_splits_legacy_text_from_latin_text(self):
        self.assertEqual(
            split_bamini_segments("அரசு Scheme"),
            [("muR", True), (" Scheme", False)],
        )


if __name__ == "__main__":
    unittest.main()
