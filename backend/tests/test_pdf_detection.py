import unittest
from unittest.mock import patch

try:
    import pymupdf as fitz
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models.ocr import OcrTextBlock
    from app.services.pdf_detection import detect_pdf_content
except ImportError:
    fitz = None
    TestClient = None
    app = None
    OcrTextBlock = None
    detect_pdf_content = None


def _scanned_page(document):
    page = document.new_page(width=300, height=400)
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), 0)
    pixmap.clear_with(230)
    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
    return page


class StubOcrProvider:
    def __init__(self, confidence=95):
        self.confidence = confidence

    def validate_language(self, language):
        return None

    def recognize(self, image, language, page_number):
        return [OcrTextBlock(text="OCR paragraph", confidence=self.confidence,
                             bbox=(100, 150, 900, 220), page=page_number)]


class StubTranslationProvider:
    async def translate(self, text, source_language, target_language):
        return f"[{target_language}] {text}"


@unittest.skipIf(fitz is None, "PyMuPDF is required for PDF detection tests")
class PdfDetectionTests(unittest.TestCase):
    def test_health_reports_translation_provider_for_privacy_notice(self):
        response = TestClient(app).get("/health")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")
        self.assertIsInstance(response.json()["translation_provider"], str)

    def test_pdf_detection_endpoint_reports_scanned_document(self):
        document = fitz.open()
        _scanned_page(document)
        pdf_bytes = document.tobytes()
        document.close()

        response = TestClient(app).post(
            "/pdf/detect",
            files={"file": ("scan.pdf", pdf_bytes, "application/pdf")},
        )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["content_type"], "scanned")
        self.assertEqual(response.json()["scanned_pages"], 1)

    def test_pdf_detection_endpoint_enforces_page_limit(self):
        document = fitz.open()
        document.new_page()
        document.new_page()
        pdf_bytes = document.tobytes()
        document.close()

        with patch("app.utils.pdf_limits.MAX_PAGES", 1):
            response = TestClient(app).post(
                "/pdf/detect",
                files={"file": ("two-pages.pdf", pdf_bytes, "application/pdf")},
            )

        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["detail"], "This PDF has 2 pages; the current limit is 1 page.")

    def test_detects_text_based_pdf(self):
        document = fitz.open()
        document.new_page().insert_text((50, 70), "Selectable text in this page.")

        result = detect_pdf_content(document)

        self.assertEqual(result.content_type, "text")
        self.assertEqual(result.text_pages, 1)
        document.close()

    def test_detects_scanned_pdf(self):
        document = fitz.open()
        _scanned_page(document)

        result = detect_pdf_content(document)

        self.assertEqual(result.content_type, "scanned")
        self.assertEqual(result.scanned_pages, 1)
        document.close()

    def test_detects_mixed_pdf(self):
        document = fitz.open()
        document.new_page().insert_text((50, 70), "Selectable text in this page.")
        _scanned_page(document)

        result = detect_pdf_content(document)

        self.assertEqual(result.content_type, "mixed")
        self.assertEqual(result.text_pages, 1)
        self.assertEqual(result.scanned_pages, 1)
        document.close()

    def test_detects_empty_pdf(self):
        document = fitz.open()
        document.new_page()

        result = detect_pdf_content(document)

        self.assertEqual(result.content_type, "empty")
        self.assertEqual(result.empty_pages, 1)
        document.close()

    def test_scanned_pdf_endpoint_uses_ocr_and_returns_docx(self):
        document = fitz.open()
        _scanned_page(document)
        pdf_bytes = document.tobytes()
        document.close()

        with patch("app.routes.translations.create_translation_provider", return_value=StubTranslationProvider()), \
             patch("app.routes.translations.create_ocr_provider", return_value=StubOcrProvider()):
            response = TestClient(app).post(
                "/translate/pdf",
                data={"source_language": "en", "target_language": "ta"},
                files={"file": ("scan.pdf", pdf_bytes, "application/pdf")},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("wordprocessingml.document", response.headers["content-type"])
        from docx import Document
        import io
        translated = Document(io.BytesIO(response.content))
        self.assertIn("[ta] OCR paragraph", translated.paragraphs[0].text)
        self.assertFalse(translated.paragraphs[0]._p.xpath(".//w:drawing"))

    def test_low_confidence_ocr_is_rejected(self):
        document = fitz.open()
        _scanned_page(document)
        pdf_bytes = document.tobytes()
        document.close()

        with patch("app.routes.translations.create_translation_provider", return_value=StubTranslationProvider()), \
             patch("app.routes.translations.create_ocr_provider", return_value=StubOcrProvider(confidence=10)):
            response = TestClient(app).post(
                "/translate/pdf",
                data={"source_language": "en", "target_language": "ta"},
                files={"file": ("scan.pdf", pdf_bytes, "application/pdf")},
            )

        self.assertEqual(response.status_code, 422)
        self.assertIn("recognize text confidently", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
