import asyncio
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.services.docx_translation import translate_docx


class RecordingProvider:
    def __init__(self):
        self.calls = []

    async def translate(self, text, source_language, target_language):
        self.calls.append((text, source_language, target_language))
        return f"[{target_language}] {text}"


class DocxTranslationTests(unittest.TestCase):
    def test_translates_paragraphs_tables_and_headers_with_cache(self):
        document = Document()
        document.add_heading("தமிழ்நாடு அரசு", level=1)
        document.add_paragraph("தமிழ்நாடு அரசு புதிய திட்டத்தை அறிமுகப்படுத்தியுள்ளது.")
        document.add_paragraph("")
        document.add_paragraph("தமிழ்நாடு அரசு புதிய திட்டத்தை அறிமுகப்படுத்தியுள்ளது.")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "தமிழ்நாடு அரசு"
        document.sections[0].header.paragraphs[0].text = "அரசு தலைப்பு"

        provider = RecordingProvider()
        translated_count = asyncio.run(translate_docx(document, provider, "ta", "en"))

        self.assertEqual(translated_count, 5)
        self.assertEqual(len(provider.calls), 3)
        self.assertTrue(all(call[1:] == ("ta", "en") for call in provider.calls))
        self.assertTrue(document.paragraphs[0].text.startswith("[en]"))
        self.assertEqual(document.paragraphs[2].text, "")
        self.assertTrue(document.tables[0].cell(0, 0).text.startswith("[en]"))
        self.assertTrue(document.sections[0].header.paragraphs[0].text.startswith("[en]"))

    def test_supports_english_to_tamil_and_preserves_urls(self):
        document = Document()
        paragraph = document.add_paragraph("Read this page: https://example.com/guide")
        provider = RecordingProvider()

        asyncio.run(translate_docx(document, provider, "en", "ta"))

        self.assertEqual(provider.calls[0][1:], ("en", "ta"))
        self.assertIn("https://example.com/guide", paragraph.text)
        self.assertNotIn("THODARURLTOKEN", paragraph.text)

    def test_skips_url_only_code_and_page_number_paragraphs(self):
        document = Document()
        document.add_paragraph("https://example.com")
        document.add_paragraph("123")
        document.add_paragraph("import os")
        provider = RecordingProvider()

        count = asyncio.run(translate_docx(document, provider, "en", "ta"))

        self.assertEqual(count, 0)
        self.assertEqual(provider.calls, [])

    def test_preserves_heading_inline_styles_and_paragraph_formatting(self):
        document = Document()
        paragraph = document.add_heading(level=2)
        paragraph.add_run("A formatted ")
        bold_run = paragraph.add_run("bold phrase")
        bold_run.bold = True
        italic_run = paragraph.add_run(" plus italic")
        italic_run.italic = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.keep_with_next = True
        provider = RecordingProvider()

        asyncio.run(translate_docx(document, provider, "en", "ta"))

        self.assertEqual(paragraph.text, "[ta] A formatted bold phrase plus italic")
        self.assertEqual(paragraph.style.name, "Heading 2")
        self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(paragraph.paragraph_format.space_before, Pt(6))
        self.assertEqual(paragraph.paragraph_format.space_after, Pt(12))
        self.assertTrue(paragraph.paragraph_format.keep_with_next)
        self.assertTrue(bold_run.bold)
        self.assertTrue(italic_run.italic)
        self.assertTrue(bold_run.text)
        self.assertTrue(italic_run.text)

    def test_preserves_mixed_formatting_in_table_cells(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        paragraph = cell.paragraphs[0]
        paragraph.add_run("Source ")
        emphasized = paragraph.add_run("label")
        emphasized.bold = True
        emphasized.font.name = "Arial"
        emphasized.font.size = Pt(14)
        paragraph.paragraph_format.space_after = Pt(8)
        provider = RecordingProvider()

        asyncio.run(translate_docx(document, provider, "en", "ta"))

        self.assertEqual(paragraph.text, "[ta] Source label")
        self.assertTrue(emphasized.bold)
        self.assertEqual(emphasized.font.name, "Arial")
        self.assertEqual(emphasized.font.size, Pt(14))
        self.assertEqual(paragraph.paragraph_format.space_after, Pt(8))


if __name__ == "__main__":
    unittest.main()
