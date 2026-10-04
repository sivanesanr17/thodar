import unittest

from docx import Document

from app.services.docx_translation import _replace_paragraph_text


class UnicodeTamilOutputTests(unittest.TestCase):
    def test_keeps_tamil_as_unicode_and_preserves_run_formatting(self):
        document = Document()
        paragraph = document.add_paragraph()
        source_run = paragraph.add_run("source")
        source_run.bold = True
        tamil = "\u0ba4\u0bae\u0bbf\u0bb4\u0bcd\u0ba8\u0bbe\u0b9f\u0bc1"

        _replace_paragraph_text(paragraph, tamil)

        self.assertEqual(paragraph.text, tamil)
        self.assertTrue(source_run.bold)
        self.assertTrue(any("\u0b80" <= char <= "\u0bff" for char in paragraph.text))
        self.assertNotEqual(source_run.font.name, "Bamini")


if __name__ == "__main__":
    unittest.main()
